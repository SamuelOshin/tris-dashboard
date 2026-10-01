"""
Ticket 5 — ERP/BOM mapping engine.

Required by the ticket:
- test_mapping_profile_save_and_reload_round_trip
- test_malformed_source_file_produces_readable_row_errors
- test_circuit_breaker_trips_on_badly_mapped_column
The rest cover Generic / SAP-style / Dynamics-style imports end to end, duplicates, referential
integrity, defaults, Excel input, the error-log download, server-side roles and UI copy.
"""

import json
from io import BytesIO
from pathlib import Path

import pytest
from httpx import AsyncClient
from openpyxl import Workbook
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.modules.v1.ingestion.models.ingestion_job import IngestionJob
from app.api.modules.v1.manufacturing.models import (
    BOMEntry,
    Material,
    PurchaseRecord,
)
from app.api.modules.v1.manufacturing.service.mapping_engine import _flush
from app.api.modules.v1.manufacturing.service.mapping_service import _safe_cell
from app.api.modules.v1.manufacturing.service.mapping_targets import FieldSpec
from app.api.modules.v1.manufacturing.service.value_coercion import FieldValueError, coerce_value
from app.api.modules.v1.suppliers.models.supplier import Supplier
from tests.conftest import make_principal

BASE = "/api/v1/manufacturing/mapping"
REPO_ROOT = Path(__file__).resolve().parents[4]

GENERIC_PURCHASES = (
    "po,material,vendor,order_date,qty,price\n"
    "PO-1,MAT-1,SUP-M1,2026-01-05,100,9.50\n"
    "PO-2,MAT-1,SUP-M1,2026-02-05,120,9.80\n"
)
GENERIC_PURCHASE_MAPPING = {
    "purchase_reference": "po",
    "material_id": "material",
    "supplier_id": "vendor",
    "purchase_date": "order_date",
    "quantity": "qty",
    "unit_price": "price",
}
SAP_MATERIALS = "MATNR;MAKTX;MATKL;MEINS\n000000000000100234;Aluminium sheet;Metals;KG\n"
SAP_MATERIAL_MAPPING = {
    "material_id": "MATNR",
    "description": "MAKTX",
    "category": "MATKL",
    "unit_of_measure": "MEINS",
}
SAP_PURCHASES = (
    "EBELN;MATNR;LIFNR;BEDAT;MENGE;NETPR;WAERS\n"
    "4500000001;000000000000100234;SUP-M1;20260105;1.250,00;9,5;eur\n"
)


async def _seed_base(session: AsyncSession) -> None:
    session.add(Material(material_id="MAT-1", description="Steel coil", unit_of_measure="kg"))
    session.add(
        Supplier(
            supplier_id="SUP-M1",
            name="Metals Co",
            category="Metals",
            risk_tier="Low",
            bank_account="123456789012",
            routing_number="123456789",
            status="Active",
        )
    )
    await session.commit()


def _config(target: str, mapping: dict, **extra) -> str:
    return json.dumps({"target": target, "field_mapping": mapping, **extra})


async def _post(client: AsyncClient, route: str, content, name: str, config: str):
    data = content if isinstance(content, bytes) else content.encode("utf-8")
    return await client.post(
        f"{BASE}/{route}", files={"file": (name, data, "text/csv")}, data={"config": config}
    )


async def _import(client: AsyncClient, content, target: str, mapping: dict, **extra):
    return await _post(client, "import", content, "data.csv", _config(target, mapping, **extra))


def _numbered_purchases(good: int, bad_quantity: int = 0) -> str:
    lines = ["po,material,vendor,order_date,qty,price"]
    for i in range(good):
        lines.append(f"PO-{i},MAT-1,SUP-M1,2026-01-05,{10 + i},9.5")
    for i in range(bad_quantity):
        lines.append(f"PO-BAD-{i},MAT-1,SUP-M1,2026-01-05,lots,9.5")
    return "\n".join(lines) + "\n"


async def _count(session: AsyncSession, model) -> int:
    return (await session.execute(select(func.count()).select_from(model))).scalar_one()


# ── Required: profile save / reload ──────────────────────────────────────────


@pytest.mark.asyncio
async def test_mapping_profile_save_and_reload_round_trip(async_client: AsyncClient):
    payload = {
        "name": "Plant A purchases",
        "description": "Weekly PO export",
        "source_profile": "sap_style",
        "target": "purchase_records",
        "field_mapping": {
            "material_id": "MATNR",
            "quantity": "MENGE",
            "unit_price": "NETPR",
            "purchase_date": "BEDAT",
        },
        "defaults": {"currency": "EUR"},
    }
    saved = await async_client.post(f"{BASE}/profiles", json=payload)
    assert saved.status_code == 201
    profile_id = saved.json()["data"]["profile_id"]

    reloaded = (await async_client.get(f"{BASE}/profiles/{profile_id}")).json()["data"]
    for key in ("name", "description", "source_profile", "target", "field_mapping", "defaults"):
        assert reloaded[key] == payload[key]
    assert reloaded["created_by"] == "USR-TEST-001"

    listed = (await async_client.get(f"{BASE}/profiles?target=purchase_records")).json()["data"]
    assert [p["profile_id"] for p in listed["profiles"]] == [profile_id]
    assert (await async_client.get(f"{BASE}/profiles?target=materials")).json()["data"][
        "profiles"
    ] == []

    again = await async_client.post(f"{BASE}/profiles", json=payload)
    assert again.status_code == 409  # names are unique

    assert (await async_client.delete(f"{BASE}/profiles/{profile_id}")).status_code == 200
    assert (await async_client.get(f"{BASE}/profiles/{profile_id}")).status_code == 404


@pytest.mark.asyncio
async def test_saved_profile_drives_a_later_import(
    async_client: AsyncClient, db_session: AsyncSession
):
    await _seed_base(db_session)
    saved = await async_client.post(
        f"{BASE}/profiles",
        json={
            "name": "Generic POs",
            "target": "purchase_records",
            "field_mapping": GENERIC_PURCHASE_MAPPING,
        },
    )
    profile = saved.json()["data"]
    res = await _import(
        async_client,
        GENERIC_PURCHASES,
        "purchase_records",
        profile["field_mapping"],
        profile_id=profile["profile_id"],
    )
    summary = res.json()["data"]["summary"]
    assert summary["profile_id"] == profile["profile_id"]
    assert summary["rows_accepted"] == 2

    missing = await _import(
        async_client,
        GENERIC_PURCHASES,
        "purchase_records",
        GENERIC_PURCHASE_MAPPING,
        profile_id="PROF-nope",
    )
    assert missing.status_code == 404


@pytest.mark.asyncio
async def test_profile_definition_is_validated(async_client: AsyncClient):
    base = {"name": "x", "target": "purchase_records", "field_mapping": {"material_id": "M"}}
    cases = {
        "required fields missing": base,
        "unknown field": {**base, "field_mapping": {**GENERIC_PURCHASE_MAPPING, "colour": "c"}},
        "bad default": {
            **base,
            "field_mapping": GENERIC_PURCHASE_MAPPING,
            "defaults": {"quantity": "lots"},
        },
        "unknown target": {**base, "target": "nope"},
    }
    for label, body in cases.items():
        res = await async_client.post(f"{BASE}/profiles", json=body)
        assert res.status_code == 422, label


# ── Targets and preview ──────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_targets_distinguish_required_and_optional_fields(async_client: AsyncClient):
    data = (await async_client.get(f"{BASE}/targets")).json()["data"]
    assert len(data["targets"]) == 11
    purchases = next(t for t in data["targets"] if t["key"] == "purchase_records")
    by_name = {f["name"]: f for f in purchases["fields"]}
    assert by_name["quantity"]["required"] is True
    assert by_name["purchase_reference"]["required"] is False
    labels = {p["key"]: p["label"] for p in data["source_profiles"]}
    assert labels["sap_style"] == "SAP-style import demonstration"
    assert labels["dynamics_style"] == "Dynamics 365-style import demonstration"


@pytest.mark.asyncio
async def test_preview_shows_columns_samples_and_suggests_the_matching_style(
    async_client: AsyncClient,
):
    files = {"file": ("po.csv", SAP_PURCHASES.encode(), "text/csv")}
    res = await async_client.post(
        f"{BASE}/preview", files=files, data={"target": "purchase_records"}
    )
    data = res.json()["data"]
    assert data["columns"] == ["EBELN", "MATNR", "LIFNR", "BEDAT", "MENGE", "NETPR", "WAERS"]
    assert data["sample_rows"][0]["MATNR"] == "000000000000100234"  # leading zeros kept
    assert data["total_rows"] == 1
    assert data["suggested_profile"] == "sap_style"
    assert data["suggestions"]["sap_style"]["quantity"] == "MENGE"

    dyn = (
        "PurchId,ItemNumber,VendAccount,OrderDate,PurchQty,PurchPrice,CurrencyCode\n"
        "P1,M1,V1,2026-01-01,1,2\n"
    )
    files = {"file": ("dyn.csv", dyn.encode(), "text/csv")}
    res = await async_client.post(
        f"{BASE}/preview", files=files, data={"target": "purchase_records"}
    )
    assert res.json()["data"]["suggested_profile"] == "dynamics_style"


# ── End-to-end imports: Generic, SAP-style, Dynamics-style ───────────────────


@pytest.mark.asyncio
async def test_generic_import_end_to_end(async_client: AsyncClient, db_session: AsyncSession):
    await _seed_base(db_session)
    res = await _import(
        async_client,
        GENERIC_PURCHASES,
        "purchase_records",
        GENERIC_PURCHASE_MAPPING,
        dataset_id="SYNTHETIC-A",
    )
    assert res.status_code == 200
    body = res.json()["data"]
    assert body["status"] == "COMPLETED"
    assert body["summary"]["rows_accepted"] == 2
    assert body["summary"]["rows_rejected"] == 0
    assert body["summary"]["source_profile_label"] == "Generic CSV / Excel"

    rows = (
        (await db_session.execute(select(PurchaseRecord).order_by(PurchaseRecord.id)))
        .scalars()
        .all()
    )
    assert [(r.purchase_reference, r.quantity, r.unit_price) for r in rows] == [
        ("PO-1", 100.0, 9.5),
        ("PO-2", 120.0, 9.8),
    ]
    assert rows[0].dataset_id == "SYNTHETIC-A"
    assert rows[0].currency == "USD"  # default, reported as a warning
    codes = {(w["code"], w["field"]) for w in body["summary"]["warnings"]}
    assert ("default_applied", "currency") in codes

    job = await db_session.get(IngestionJob, body["job_id"])
    assert (job.status, job.inserted_rows, job.uploaded_by) == ("COMPLETED", 2, "USR-TEST-001")
    assert job.summary_report["field_mapping"] == GENERIC_PURCHASE_MAPPING


@pytest.mark.asyncio
async def test_sap_style_import_end_to_end(async_client: AsyncClient, db_session: AsyncSession):
    db_session.add(
        Supplier(
            supplier_id="SUP-M1",
            name="Metals Co",
            category="Metals",
            risk_tier="Low",
            status="Active",
        )
    )
    await db_session.commit()

    masters = await _import(
        async_client, SAP_MATERIALS, "materials", SAP_MATERIAL_MAPPING, source_profile="sap_style"
    )
    assert masters.json()["data"]["summary"]["rows_accepted"] == 1
    material = await db_session.get(Material, "000000000000100234")
    assert material is not None and material.unit_of_measure == "KG"

    # European number format is not guessed at: it must be reported, not silently misread.
    mapping = {
        "purchase_reference": "EBELN",
        "material_id": "MATNR",
        "supplier_id": "LIFNR",
        "purchase_date": "BEDAT",
        "quantity": "MENGE",
        "unit_price": "NETPR",
        "currency": "WAERS",
    }
    bad = (
        await _import(
            async_client, SAP_PURCHASES, "purchase_records", mapping, source_profile="sap_style"
        )
    ).json()["data"]
    assert bad["summary"]["rows_rejected"] == 1
    assert bad["error_log"][0]["field"] == "quantity"

    good_file = SAP_PURCHASES.replace("1.250,00", "1250.00").replace("9,5", "9.5")
    ok = (
        await _import(
            async_client, good_file, "purchase_records", mapping, source_profile="sap_style"
        )
    ).json()["data"]
    assert ok["summary"]["rows_accepted"] == 1
    record = (await db_session.execute(select(PurchaseRecord))).scalars().one()
    assert record.material_id == "000000000000100234"  # leading zeros preserved
    assert str(record.purchase_date) == "2026-01-05"  # YYYYMMDD understood
    assert record.currency == "EUR"
    assert ok["summary"]["source_profile_label"] == "SAP-style import demonstration"


@pytest.mark.asyncio
async def test_dynamics_style_bom_import(async_client: AsyncClient, db_session: AsyncSession):
    await _seed_base(db_session)
    bom = (
        "ProductNumber,ItemNumber,BOMQuantity,UnitId,FromDate\n"
        "SKU-9,MAT-1,2.5,kg,2026-01-01\nSKU-9,MAT-1,3,kg,2026-06-01\n"
    )
    files = {"file": ("bom.csv", bom.encode(), "text/csv")}
    preview = (
        await async_client.post(f"{BASE}/preview", files=files, data={"target": "bom_entries"})
    ).json()["data"]
    assert preview["suggested_profile"] == "dynamics_style"
    suggestion = preview["suggestions"]["dynamics_style"]
    assert (
        suggestion["product_sku"] == "ProductNumber" and suggestion["material_id"] == "ItemNumber"
    )

    res = await _import(
        async_client, bom, "bom_entries", suggestion, source_profile="dynamics_style"
    )
    assert res.json()["data"]["summary"]["rows_accepted"] == 2
    assert await _count(db_session, BOMEntry) == 2


@pytest.mark.asyncio
async def test_excel_workbook_import_with_sheet_selection(
    async_client: AsyncClient, db_session: AsyncSession
):
    await _seed_base(db_session)
    book = Workbook()
    book.active.title = "Notes"
    book.active.append(["nothing", "here"])
    sheet = book.create_sheet("Purchases")
    sheet.append(["po", "material", "vendor", "order_date", "qty", "price"])
    sheet.append(["PO-X", "MAT-1", "SUP-M1", "2026-03-01", 40, 9.1])
    buffer = BytesIO()
    book.save(buffer)

    files = {"file": ("wb.xlsx", buffer.getvalue(), "application/octet-stream")}
    preview = (
        await async_client.post(
            f"{BASE}/preview",
            files=files,
            data={"target": "purchase_records", "sheet": "Purchases"},
        )
    ).json()["data"]
    assert preview["sheets"] == ["Notes", "Purchases"] and preview["total_rows"] == 1

    res = await async_client.post(
        f"{BASE}/import",
        files={"file": ("wb.xlsx", buffer.getvalue(), "application/octet-stream")},
        data={"config": _config("purchase_records", GENERIC_PURCHASE_MAPPING, sheet="Purchases")},
    )
    assert res.json()["data"]["summary"]["rows_accepted"] == 1


# ── Required: malformed file gives a readable error log ──────────────────────


@pytest.mark.asyncio
async def test_malformed_source_file_produces_readable_row_errors(
    async_client: AsyncClient, db_session: AsyncSession
):
    await _seed_base(db_session)
    lines = _numbered_purchases(40).splitlines()
    lines.insert(3, "PO-RAGGED,MAT-1")  # too few cells
    lines.insert(6, "PO-DATE,MAT-1,SUP-M1,31/31/2026,5,9.5")  # impossible date
    lines.insert(9, "PO-NEG,MAT-1,SUP-M1,2026-01-05,-4,9.5")  # negative quantity
    lines.insert(12, "PO-GHOST,MAT-NOPE,SUP-M1,2026-01-05,4,9.5")  # unknown material
    res = await _import(
        async_client, "\n".join(lines) + "\n", "purchase_records", GENERIC_PURCHASE_MAPPING
    )

    assert res.status_code == 200  # an error log, not a crash
    body = res.json()["data"]
    assert body["status"] == "COMPLETED_WITH_ERRORS"
    assert body["summary"]["rows_accepted"] == 40
    assert body["summary"]["rows_rejected"] == 4
    log = {e["field"]: e for e in body["error_log"]}
    assert "2 cells but the header has 6" in log["file_structure"]["error"]
    assert "not a recognised date" in log["purchase_date"]["error"]
    assert "greater than 0" in log["quantity"]["error"]
    assert "Material 'MAT-NOPE' does not exist" in log["material_id"]["error"]
    assert all(isinstance(e["row"], int) and e["row"] > 1 for e in body["error_log"])
    assert log["quantity"]["raw_value"]["qty"] == "-4"  # the offending source row is shown
    assert await _count(db_session, PurchaseRecord) == 40


@pytest.mark.asyncio
async def test_unreadable_files_give_readable_errors(async_client: AsyncClient):
    cfg = _config("purchase_records", GENERIC_PURCHASE_MAPPING)
    cases = [
        ("binary.csv", b"\x00\x01\x02\x03" * 50, "not a text CSV"),
        ("empty.csv", b"", "empty"),
        ("headers_only.csv", b"po,material\n", "no data rows"),
        ("old.xls", b"PK", "Unsupported file"),
        ("fake.xlsx", b"this is not a workbook", "Excel workbook"),
        ("dup.csv", b"a,a\n1,2\n", "repeats these column names"),
    ]
    for name, content, expected in cases:
        res = await _post(async_client, "validate", content, name, cfg)
        assert res.status_code == 422, name
        assert expected in res.json()["message"], (name, res.json()["message"])


# ── Required: circuit breaker ────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_circuit_breaker_trips_on_badly_mapped_column(
    async_client: AsyncClient, db_session: AsyncSession
):
    await _seed_base(db_session)
    # "vendor" holds text, so mapping it to quantity makes every row fail.
    wrong = {**GENERIC_PURCHASE_MAPPING, "quantity": "vendor"}
    res = await _import(async_client, _numbered_purchases(30), "purchase_records", wrong)
    body = res.json()["data"]

    assert body["status"] == "FAILED"
    breaker = body["summary"]["circuit_breaker"]
    assert breaker["tripped"] is True and breaker["stage"] == "validation"
    assert breaker["ratio"] == 1.0 and breaker["threshold"] == 0.2
    assert breaker["suspect_field"] == "quantity" and breaker["suspect_column"] == "vendor"
    assert "mapped from column 'vendor'" in breaker["hint"]
    assert body["summary"]["rows_accepted"] == 0
    assert await _count(db_session, PurchaseRecord) == 0  # nothing half-loaded

    job = await db_session.get(IngestionJob, body["job_id"])
    assert job.status == "FAILED"
    assert any(e["field"] == "circuit_breaker" for e in job.error_log)


@pytest.mark.asyncio
async def test_circuit_breaker_thresholds_match_v14_semantics(
    async_client: AsyncClient, db_session: AsyncSession
):
    await _seed_base(db_session)
    # exactly 20% bad is tolerated (the breaker trips only above 20%)
    at_limit = (
        await _import(
            async_client, _numbered_purchases(8, 2), "purchase_records", GENERIC_PURCHASE_MAPPING
        )
    ).json()["data"]
    assert at_limit["status"] == "COMPLETED_WITH_ERRORS"
    assert at_limit["summary"]["circuit_breaker"] is None
    assert at_limit["summary"]["rows_accepted"] == 8

    # 3 of 10 is over the limit
    over = (
        await _import(
            async_client,
            _numbered_purchases(7, 3).replace("PO-", "Q-"),
            "purchase_records",
            GENERIC_PURCHASE_MAPPING,
        )
    ).json()["data"]
    assert over["status"] == "FAILED"

    # fewer than 10 rows is too small a sample to trip, even if every row is bad
    tiny = (
        await _import(
            async_client, _numbered_purchases(0, 4), "purchase_records", GENERIC_PURCHASE_MAPPING
        )
    ).json()["data"]
    assert tiny["status"] == "COMPLETED_WITH_ERRORS" and tiny["summary"]["circuit_breaker"] is None


# ── Validation, duplicates, references, defaults ─────────────────────────────


@pytest.mark.asyncio
async def test_validate_is_a_dry_run(async_client: AsyncClient, db_session: AsyncSession):
    await _seed_base(db_session)
    res = await _post(
        async_client,
        "validate",
        GENERIC_PURCHASES,
        "po.csv",
        _config("purchase_records", GENERIC_PURCHASE_MAPPING),
    )
    body = res.json()["data"]
    assert body["status"] == "VALIDATED" and body["job_id"] is None
    assert body["summary"]["dry_run"] is True and body["summary"]["rows_accepted"] == 2
    assert await _count(db_session, PurchaseRecord) == 0
    assert await _count(db_session, IngestionJob) == 0


@pytest.mark.asyncio
async def test_required_fields_and_columns_are_checked_before_any_row(async_client: AsyncClient):
    short = {k: v for k, v in GENERIC_PURCHASE_MAPPING.items() if k != "quantity"}
    res = await _import(async_client, GENERIC_PURCHASES, "purchase_records", short)
    assert res.status_code == 422 and "quantity" in res.json()["message"]

    ghost = {**GENERIC_PURCHASE_MAPPING, "quantity": "no_such_column"}
    res = await _import(async_client, GENERIC_PURCHASES, "purchase_records", ghost)
    assert res.status_code == 422 and "no_such_column" in res.json()["message"]

    res = await _post(async_client, "import", GENERIC_PURCHASES, "po.csv", "not json")
    assert res.status_code == 422


@pytest.mark.asyncio
async def test_duplicates_are_skipped_or_rejected_by_strategy(
    async_client: AsyncClient, db_session: AsyncSession
):
    await _seed_base(db_session)
    first = await _import(
        async_client, GENERIC_PURCHASES, "purchase_records", GENERIC_PURCHASE_MAPPING
    )
    assert first.json()["data"]["summary"]["rows_accepted"] == 2

    again = (
        await _import(async_client, GENERIC_PURCHASES, "purchase_records", GENERIC_PURCHASE_MAPPING)
    ).json()["data"]
    assert again["summary"]["rows_duplicate"] == 2 and again["summary"]["rows_accepted"] == 0
    assert again["status"] == "COMPLETED"  # skipping a duplicate is not an error

    strict = (
        await _import(
            async_client,
            GENERIC_PURCHASES,
            "purchase_records",
            GENERIC_PURCHASE_MAPPING,
            duplicate_strategy="fail",
        )
    ).json()["data"]
    assert strict["summary"]["rows_rejected"] == 2
    assert "already stored" in strict["error_log"][0]["error"]

    in_file = GENERIC_PURCHASES + "PO-2,MAT-1,SUP-M1,2026-02-05,120,9.80\n"
    dup = (
        await _import(
            async_client, in_file.replace("PO-", "N-"), "purchase_records", GENERIC_PURCHASE_MAPPING
        )
    ).json()["data"]
    assert dup["summary"]["rows_accepted"] == 2 and dup["summary"]["rows_duplicate"] == 1
    assert await _count(db_session, PurchaseRecord) == 4


@pytest.mark.asyncio
async def test_references_to_unknown_suppliers_are_rejected_readably(
    async_client: AsyncClient, db_session: AsyncSession
):
    await _seed_base(db_session)
    text_ = _numbered_purchases(20) + "PO-S,MAT-1,SUP-GHOST,2026-01-05,5,9.5\n"
    body = (
        await _import(async_client, text_, "purchase_records", GENERIC_PURCHASE_MAPPING)
    ).json()["data"]
    entry = body["error_log"][0]
    assert entry["field"] == "supplier_id"
    assert "Supplier 'SUP-GHOST' does not exist" in entry["error"]


@pytest.mark.asyncio
async def test_defaults_missing_values_and_unmapped_fields_are_reported(
    async_client: AsyncClient,
):
    csv_ = "id,name,group\nM-1,Copper wire,\nM-2,Zinc ingot,Metals\n"
    res = await _import(
        async_client,
        csv_,
        "materials",
        {"material_id": "id", "description": "name", "category": "group"},
        defaults={"unit_of_measure": "kg"},
    )
    summary = res.json()["data"]["summary"]
    assert summary["rows_accepted"] == 2
    assert summary["missing_values"] == {"category": 1}
    assert "is_active" in summary["unmapped_optional_fields"]
    assert summary["defaults"] == {"unit_of_measure": "kg"}


@pytest.mark.asyncio
async def test_cross_field_rules(async_client: AsyncClient, db_session: AsyncSession):
    await _seed_base(db_session)
    rows = "\n".join(f"MAT-1,2026-01-{d:02d},2026-01-{d - 1:02d},10" for d in range(5, 15))
    csv_ = "m,start,end,std\n" + rows + "\nMAT-1,2026-02-01,2026-02-28,\n"
    body = (
        await _import(
            async_client,
            csv_,
            "variance_inputs",
            {
                "material_id": "m",
                "period_start": "start",
                "period_end": "end",
                "standard_price": "std",
            },
        )
    ).json()["data"]
    assert body["status"] == "FAILED"  # 10 of 11 rows end before they start
    assert body["summary"]["circuit_breaker"]["suspect_field"] == "period_end"


# ── Error-log download, roles and housekeeping ───────────────────────────────


@pytest.mark.asyncio
async def test_error_log_csv_download_is_escaped(
    async_client: AsyncClient, db_session: AsyncSession
):
    await _seed_base(db_session)
    csv_ = _numbered_purchases(30) + "PO-F,=HYPERLINK(1),SUP-M1,2026-01-05,5,9.5\n"
    job_id = (
        await _import(async_client, csv_, "purchase_records", GENERIC_PURCHASE_MAPPING)
    ).json()["data"]["job_id"]
    res = await async_client.get(f"{BASE}/jobs/{job_id}/errors.csv")
    assert res.status_code == 200 and res.headers["content-type"].startswith("text/csv")
    lines = res.text.strip().splitlines()
    assert lines[0] == "row,sheet,field,error,source_values"
    assert len(lines) == 2 and "does not exist" in lines[1]

    assert _safe_cell("=1+1") == "'=1+1" and _safe_cell("@x") == "'@x" and _safe_cell("ok") == "ok"
    plain = IngestionJob(job_id="INGEST-plain", uploaded_by="USR-TEST-001")
    db_session.add(plain)
    await db_session.commit()
    assert (await async_client.get(f"{BASE}/jobs/INGEST-plain/errors.csv")).status_code == 404


@pytest.mark.asyncio
async def test_server_side_roles(client_as, db_session: AsyncSession):
    await _seed_base(db_session)
    files = {"file": ("po.csv", GENERIC_PURCHASES.encode(), "text/csv")}
    body = {"name": "r", "target": "purchase_records", "field_mapping": GENERIC_PURCHASE_MAPPING}

    for role in ("read_only_reviewer", "verifier", "process_owner"):
        who = make_principal("USR-ADMIN-001", f"x_{role}", "X", role)
        async with client_as(who) as client:
            assert (await client.get(f"{BASE}/targets")).status_code == 403, role
            assert (
                await client.post(
                    f"{BASE}/preview", files=files, data={"target": "purchase_records"}
                )
            ).status_code == 403
            assert (
                await _import(
                    client, GENERIC_PURCHASES, "purchase_records", GENERIC_PURCHASE_MAPPING
                )
            ).status_code == 403
            assert (await client.get(f"{BASE}/profiles")).status_code == 403

    reviewer = make_principal("USR-ADMIN-001", "rev", "Rev", "reviewer")
    async with client_as(reviewer) as client:
        assert (await client.get(f"{BASE}/targets")).status_code == 200
        res = await _import(client, GENERIC_PURCHASES, "purchase_records", GENERIC_PURCHASE_MAPPING)
        assert res.status_code == 200 and res.json()["data"]["summary"]["rows_accepted"] == 2
        assert (await client.get(f"{BASE}/profiles")).status_code == 200
        assert (
            await client.post(f"{BASE}/profiles", json=body)
        ).status_code == 403  # config = admin


@pytest.mark.asyncio
@pytest.mark.filterwarnings("ignore:New instance")  # the clash is deliberate
async def test_database_rejections_are_isolated_and_hide_internals(db_session: AsyncSession):
    good = Material(material_id="ISO-1", description="ok", unit_of_measure="kg")
    clash_a = Material(material_id="ISO-2", description="first", unit_of_measure="kg")
    clash_b = Material(material_id="ISO-2", description="second", unit_of_measure="kg")
    clash_b._source_row_num = 7  # type: ignore[attr-defined]
    chunk = [(good, {}), (clash_a, {}), (clash_b, {"material_id": "ISO-2"})]

    inserted, errors = await _flush(db_session, chunk, "CSV")
    assert inserted == 2 and len(errors) == 1
    assert errors[0]["row"] == 7 and errors[0]["field"] == "data_integrity"
    assert "SELECT" not in errors[0]["error"] and "materials" not in errors[0]["error"]


def test_value_coercion_rules():
    num = FieldSpec("n", "float", min_value=0, min_exclusive=True)
    dat = FieldSpec("d", "date")
    flag = FieldSpec("f", "bool")
    rate = FieldSpec("r", "float", min_value=0, max_value=1, hint="use a fraction")
    assert coerce_value(num, "1,234.50")[0] == 1234.5
    assert coerce_value(dat, "20260131")[0].isoformat() == "2026-01-31"
    assert coerce_value(dat, "31.01.2026")[0].isoformat() == "2026-01-31"
    assert coerce_value(dat, "2026-01-31 00:00:00")[0].isoformat() == "2026-01-31"
    assert coerce_value(flag, "X")[0] is True and coerce_value(flag, "no")[0] is False
    assert coerce_value(num, "  ")[0] is None
    assert coerce_value(dat, "00000000")[0] is None  # ERP "no date"
    for spec, bad in (
        (num, "0"),
        (num, "abc"),
        (num, "inf"),
        (dat, "2026-13-45"),
        (flag, "maybe"),
        (rate, "92"),
    ):
        with pytest.raises(FieldValueError):
            coerce_value(spec, bad)
    with pytest.raises(FieldValueError, match="use a fraction"):
        coerce_value(rate, "92")


@pytest.mark.asyncio
async def test_mapping_profiles_table_exists_with_unique_name(db_session: AsyncSession):
    rows = await db_session.execute(
        text("SELECT indexdef FROM pg_indexes WHERE tablename = 'mapping_profiles'")
    )
    assert any("UNIQUE" in r[0] and "(name)" in r[0] for r in rows.fetchall())


def test_ui_never_claims_a_live_erp_integration():
    forbidden = ("live sap", "live dynamics")
    for folder in ("app", "components"):
        for path in (REPO_ROOT / "frontend" / folder).rglob("*.ts*"):
            lowered = path.read_text(encoding="utf-8").lower()
            assert not any(term in lowered for term in forbidden), path


# ── QA follow-up fixes (Ticket 5 review) ─────────────────────────────────────


def _workbook_bytes(rows: list[list]) -> bytes:
    book = Workbook()
    for row in rows:
        book.active.append(row)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


async def _validate_xlsx(client: AsyncClient, content: bytes, mapping: dict):
    return await client.post(
        f"{BASE}/validate",
        files={"file": ("wb.xlsx", content, "application/octet-stream")},
        data={"config": _config("materials", mapping)},
    )


@pytest.mark.asyncio
async def test_excel_headers_get_the_same_checks_as_csv(async_client: AsyncClient):
    mapping = {"material_id": "id", "description": "name", "unit_of_measure": "u"}
    repeated = _workbook_bytes([["id", "id", "u"], ["M-1", "x", "kg"]])
    res = await _validate_xlsx(async_client, repeated, mapping)
    assert res.status_code == 422 and "repeats these column names" in res.json()["message"]

    blank = _workbook_bytes([["id", None, "u"], ["M-1", "x", "kg"]])
    res = await _validate_xlsx(async_client, blank, mapping)
    assert res.status_code == 422 and "empty column name" in res.json()["message"]

    ok = _workbook_bytes([["id", "name", "u"], ["M-1", "Zinc", "kg"]])
    res = await _validate_xlsx(async_client, ok, mapping)
    assert res.status_code == 200 and res.json()["data"]["summary"]["rows_accepted"] == 1


@pytest.mark.asyncio
async def test_over_long_identifiers_are_rejected_not_shortened(async_client: AsyncClient):
    mapping = {"material_id": "id", "description": "name", "unit_of_measure": "u"}
    csv_ = f"id,name,u\n{'A' * 80},Zinc,kg\nM-2,{'D' * 600},kg\n"
    body = (await _import(async_client, csv_, "materials", mapping)).json()["data"]
    assert body["summary"]["rows_accepted"] == 1 and body["summary"]["rows_rejected"] == 1
    entry = body["error_log"][0]
    assert entry["field"] == "material_id" and "limit is 50" in entry["error"]
    # ordinary text is still shortened, with a warning
    assert ("text_truncated", "description") in {
        (w["code"], w["field"]) for w in body["summary"]["warnings"]
    }


@pytest.mark.asyncio
async def test_too_long_default_is_rejected(async_client: AsyncClient):
    mapping = {"material_id": "id", "description": "name"}
    res = await _import(
        async_client,
        "id,name\nM-1,Zinc\n",
        "materials",
        mapping,
        defaults={"unit_of_measure": "x" * 30},
    )
    assert res.status_code == 422 and "unit_of_measure" in res.json()["message"]


@pytest.mark.asyncio
async def test_na_and_dash_are_valid_identifiers_but_empty_elsewhere(
    async_client: AsyncClient, db_session: AsyncSession
):
    mapping = {
        "material_id": "id",
        "description": "name",
        "unit_of_measure": "u",
        "category": "cat",
    }
    csv_ = "id,name,u,cat\nNA,Sodium,kg,n/a\n-,Dash part,kg,-\n"
    body = (await _import(async_client, csv_, "materials", mapping)).json()["data"]
    assert body["summary"]["rows_accepted"] == 2
    stored = (await db_session.execute(select(Material).order_by(Material.material_id))).scalars()
    rows = {m.material_id: m.category for m in stored}
    assert rows == {"-": None, "NA": None}  # codes kept, placeholder category read as empty


@pytest.mark.asyncio
async def test_library_error_text_is_not_shown_to_users(async_client: AsyncClient):
    res = await _post(
        async_client,
        "validate",
        GENERIC_PURCHASES,
        "po.csv",
        _config("purchase_records", GENERIC_PURCHASE_MAPPING, duplicate_strategy="bogus"),
    )
    message = res.json()["message"]
    assert res.status_code == 422 and "duplicate_strategy" in message
    assert "pydantic" not in message.lower() and "errors." not in message

    res = await _validate_xlsx(async_client, b"PK\x03\x04garbage", {})
    assert res.status_code == 422 and "zip" not in res.json()["message"].lower()
    assert "valid .xlsx file" in res.json()["message"]


@pytest.mark.asyncio
async def test_circuit_breaker_at_insertion_stage_rolls_back_everything(
    async_client: AsyncClient, db_session: AsyncSession, monkeypatch
):
    """Rows pass validation, are written, then the database rejects too many: all undone."""
    from app.api.modules.v1.manufacturing.service import mapping_engine

    await _seed_base(db_session)
    real_flush = mapping_engine._flush

    async def flaky_flush(session, chunk, sheet):
        inserted, errors = await real_flush(session, chunk, sheet)  # really writes the rows
        fake = [
            {
                "sheet": sheet,
                "row": i,
                "field": "data_integrity",
                "error": "rejected",
                "raw_value": {},
            }
            for i in range(10)
        ]
        return inserted - 10, errors + fake

    monkeypatch.setattr(mapping_engine, "_flush", flaky_flush)
    res = await _import(
        async_client, _numbered_purchases(30), "purchase_records", GENERIC_PURCHASE_MAPPING
    )
    body = res.json()["data"]

    assert body["status"] == "FAILED"
    assert body["summary"]["circuit_breaker"]["stage"] == "insertion"
    assert body["summary"]["rows_accepted"] == 0
    assert await _count(db_session, PurchaseRecord) == 0  # the flushed rows were rolled back
    job = await db_session.get(IngestionJob, body["job_id"])
    assert job.status == "FAILED" and job.inserted_rows == 0
